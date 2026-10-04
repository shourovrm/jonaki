from jonaki_docs import WordDoc, Sheets, Deck, html_to_docx, html_to_pptx, html_tables_to_xlsx, chart_image

for theme in ("leaf", "plain", "formal"):
    doc = WordDoc(theme=theme)
    doc.title("Solar payback in Dhaka", subtitle="5 October 2026 · theme " + theme)
    doc.contents()
    doc.heading("Costs")
    doc.paragraph("A 3 kW system costs **2.4 lakh taka** and pays back in *about 6 years* [1]. See [the tariff](https://example.org/tariff) and `net_metering=True`.")
    doc.paragraph("বাংলায় একটি অনুচ্ছেদ: সৌর প্যানেলের দাম প্রতি বছর কমছে।")
    doc.bullets(["Panels are 55% of the cost", ["Mono panels cost more"], "The inverter is **20%**"])
    doc.numbered(["Measure the roof", "Get three quotes"])
    doc.callout("Larger systems pay back sooner: 4.9 years at 10 kW.", title="Key finding")
    doc.heading("Sizes", level=2)
    doc.table([["Size", "Cost (taka)", "Payback (years)"], ["3 kW", 240000, 6.1], ["5 kW", 380000, 5.4], ["10 kW", 720000, 4.9]], caption="Table 1: system sizes")
    doc.chart("bar", ["Jan", "Apr", "Jul", "Oct"], {"3 kW": [310, 420, 290, 350], "5 kW": [520, 700, 480, 590]}, title="Monthly output (kWh)", caption="Figure 1: output")
    doc.quote("Net metering changes the payback more than panel price.")
    doc.code("payback = cost / yearly_saving\nprint(round(payback, 1))")
    doc.footer("Jonaki", page_numbers=True)
    doc.save("artifacts/word-%s.docx" % theme)

sheets = Sheets(theme="leaf")
budget = sheets.sheet("Budget", [["Month", "Panels", "Inverter"], ["Jan", 120000, 30000.5], ["Feb", 95000, 21000], ["Mar", 140000, 45000]],
                      number_formats={"Panels": "#,##0", "Inverter": "#,##0.00"}, totals={"Panels": "sum", "Inverter": "sum"})
sheets.chart(budget, "bar", title="Spend by month")
sheets.sheet("বাংলা শিট", [["নাম", "মান"], ["থিসিস", 12]])
sheets.save("artifacts/workbook.xlsx")

for theme in ("leaf", "formal"):
    deck = Deck(theme=theme)
    deck.title_slide("Rooftop solar pays back in 6 years", subtitle="Dhaka, October 2026", notes="Open with the number.")
    deck.bullet_slide("Panels are over half the cost", ["Panels: **55%**", "Inverter: 20%", "Mounting and wiring: 25%", ["Labour is included"]], notes="Three installers.")
    deck.table_slide("Larger systems pay back sooner", [["Size", "Cost (taka)", "Payback (years)"], ["3 kW", 240000, 6.1], ["5 kW", 380000, 5.4], ["10 kW", 720000, 4.9]])
    deck.chart_slide("Output peaks in April", "bar", ["Jan", "Apr", "Jul", "Oct"], {"3 kW": [310, 420, 290, 350], "5 kW": [520, 700, 480, 590]})
    deck.chart_slide("Panels dominate the cost", "pie", ["Panels", "Inverter", "Other"], {"Share": [55, 20, 25]})
    deck.text_slide("Payback at 10 kW", "4.9 years")
    deck.two_column_slide("Buy or lease", ["Lower total cost", "You own the panels"], ["No upfront cost", "Provider maintains"], left_heading="Buy", right_heading="Lease")
    deck.image_slide("The same chart as a picture", chart_image("line", ["Jan", "Apr", "Jul", "Oct"], {"3 kW": [310, 420, 290, 350]}), caption="Drawn by chart_image")
    deck.bullet_slide("সৌর শক্তির পরবর্তী ধাপ", ["নেট মিটারিংয়ের আবেদন করুন", "Get three quotes"])
    deck.save("artifacts/deck-%s.pptx" % theme)

html_to_docx("artifacts/report.html", "artifacts/from-html.docx")
html_to_pptx("artifacts/slides.html", "artifacts/from-html.pptx")
html_tables_to_xlsx("artifacts/report.html", "artifacts/from-html.xlsx")

# A user's own file as the template: its styles and layouts are used.
html_to_docx("artifacts/report.html", "artifacts/from-html-template.docx", template="artifacts/word-formal.docx")

# --- Checks: each file is opened again and its content read back. ---
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pptx.util import Inches

word = Document("artifacts/word-leaf.docx")
texts = [paragraph.text for paragraph in word.paragraphs]
assert "Solar payback in Dhaka" in texts, "title missing"
assert any("বাংলায় একটি অনুচ্ছেদ" in text for text in texts), "Bangla paragraph missing"
assert [style for style in (paragraph.style.name for paragraph in word.paragraphs) if style == "List Bullet 2"], "sub-bullet missing"
data_table = [table for table in word.tables if len(table.columns) == 3][0]
assert [cell.text for cell in data_table.rows[1].cells] == ["3 kW", "240,000", "6.1"], "table numbers wrong"
assert len(word.inline_shapes) == 1, "chart picture missing"
assert any("HYPERLINK" in relation.reltype.upper() for relation in word.part.rels.values()), "link missing"

book = load_workbook("artifacts/workbook.xlsx")
budget_sheet = book["Budget"]
assert budget_sheet["B5"].value == "=SUM(B2:B4)", "total formula missing"
assert budget_sheet["B2"].value == 120000 and budget_sheet["B2"].number_format == "#,##0", "number or format wrong"
assert budget_sheet.freeze_panes == "A2" and budget_sheet.auto_filter.ref == "A1:C4", "freeze or filter wrong"
assert len(budget_sheet._charts) == 1, "chart missing"
assert "বাংলা শিট" in book.sheetnames, "Bangla sheet name missing"

slides = Presentation("artifacts/deck-leaf.pptx")
assert len(slides.slides) == 9, "slide count"
assert slides.slide_width == Inches(13.333), "not 16:9"
assert slides.slides[0].notes_slide.notes_text_frame.text == "Open with the number.", "notes missing"
assert sum(1 for slide in slides.slides for shape in slide.shapes if shape.has_chart) == 2, "native charts missing"
assert sum(1 for slide in slides.slides for shape in slide.shapes if shape.has_table) == 1, "table missing"

from_html = Document("artifacts/from-html.docx")
html_texts = [paragraph.text for paragraph in from_html.paragraphs]
assert html_texts[0] == "Rooftop solar in Dhaka", "h1 did not become the title"
assert "Costs" in html_texts and "Sources" in html_texts, "headings missing"
assert not any("new Chart" in text for text in html_texts), "script text leaked into the document"
assert len(from_html.inline_shapes) == 1, "data-chart was not drawn"
assert [cell.text for cell in from_html.tables[0].rows[0].cells] == ["Size", "Cost (taka)", "Payback (years)"], "html table header"

html_deck = Presentation("artifacts/from-html.pptx")
assert len(html_deck.slides) == 5, "html slide count"
assert html_deck.slides[1].notes_slide.notes_text_frame.text == "Costs from three installers.", "html notes"
assert any(shape.has_chart for shape in html_deck.slides[3].shapes), "html chart not native"

html_book = load_workbook("artifacts/from-html.xlsx")
assert html_book.sheetnames == ["System sizes"], "sheet not named by its caption"
assert html_book["System sizes"]["B2"].value == 240000, "'240,000' did not become a number"

templated = Document("artifacts/from-html-template.docx")
templated_texts = [paragraph.text for paragraph in templated.paragraphs]
assert templated_texts[0] == "Rooftop solar in Dhaka", "template run lost the content"
assert "Solar payback in Dhaka" not in templated_texts, "the template's own text stayed in the new file"
assert templated.styles["Heading 1"].font.color.rgb == Document("artifacts/word-formal.docx").styles["Heading 1"].font.color.rgb, "template styles not used"

html_to_pptx("artifacts/slides.html", "artifacts/from-html-template.pptx", template="artifacts/deck-formal.pptx")
templated_deck = Presentation("artifacts/from-html-template.pptx")
assert len(templated_deck.slides) == 5, "the template's own slides stayed in the new deck"
assert templated_deck.slides[1].shapes.title.text == "Panels are over half the cost", "template layout title not filled"
print("all checks passed")
